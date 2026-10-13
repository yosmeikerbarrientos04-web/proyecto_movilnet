import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generar_excel_acta(proyecto, acta):
    """
    Genera un archivo Excel (.xlsx) que reproduce fielmente el formato
    del Acta de Recepción de Requerimiento de Solución Tecnológica (Software).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acta de Recepción"
    ws.views.sheetView[0].showGridLines = True

    # Configuración de anchos de columnas para proporción armónica tipo carta
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 24
    ws.column_dimensions['D'].width = 32

    # Bordes y rellenos
    thin_side = Side(style='thin', color='000000')
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    label_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    approved_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')

    font_title_org = Font(name='Arial', size=10, bold=True, color='000000')
    font_title_main = Font(name='Arial', size=11, bold=True, color='000000')
    font_section_header = Font(name='Arial', size=9.5, bold=True, color='000000')
    font_label = Font(name='Arial', size=8.5, bold=True, color='000000')
    font_value = Font(name='Arial', size=9, color='000000')
    font_small = Font(name='Arial', size=8, color='334155')

    def apply_style_range(start_col, start_row, end_col, end_row, font=None, fill=None, alignment=None, border=None):
        for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
            for cell in row:
                if font: cell.font = font
                if fill: cell.fill = fill
                if alignment: cell.alignment = alignment
                if border: cell.border = border

    current_row = 2

    # CABECERA
    ws.cell(row=current_row, column=1, value="GERENCIA GENERAL DE DESARROLLO DE SOLUCIONES TI")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    apply_style_range(1, current_row, 4, current_row, font=font_title_org, alignment=Alignment(horizontal='center', vertical='center'))
    current_row += 1

    ws.cell(row=current_row, column=1, value="ACTA DE RECEPCIÓN DE REQUERIMIENTO DE SOLUCIÓN TECNOLÓGICA (SOFTWARE)")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    apply_style_range(1, current_row, 4, current_row, font=font_title_main, alignment=Alignment(horizontal='center', vertical='center'))
    current_row += 2

    # FECHA
    ws.cell(row=current_row, column=3, value="FECHA:")
    ws.cell(row=current_row, column=3).font = font_label
    ws.cell(row=current_row, column=3).alignment = Alignment(horizontal='right', vertical='center')

    fecha_str = acta.fecha.strftime("%d/%m/%Y") if acta.fecha else ""
    date_cell = ws.cell(row=current_row, column=4, value=fecha_str)
    date_cell.font = font_value
    date_cell.alignment = Alignment(horizontal='center', vertical='center')
    date_cell.border = thin_border
    current_row += 2

    # SECCIÓN: REQUERIMIENTO
    ws.cell(row=current_row, column=1, value="REQUERIMIENTO")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    apply_style_range(1, current_row, 4, current_row, font=font_section_header, fill=header_fill, alignment=Alignment(horizontal='center', vertical='center'), border=thin_border)
    current_row += 1

    # Fila Proyecto / ID
    r1 = current_row
    ws.cell(row=r1, column=1, value="PROYECTO")
    ws.cell(row=r1, column=2, value=proyecto.nombre)
    ws.cell(row=r1, column=3, value="ID")
    ws.cell(row=r1, column=4, value=proyecto.identificador or "")
    current_row += 1

    # Fila Siglas / Gerencia Receptora
    r2 = current_row
    ws.cell(row=r2, column=1, value="SIGLAS DE LA SOLUCIÓN")
    ws.cell(row=r2, column=2, value=acta.siglas_solucion or "")
    ws.cell(row=r2, column=3, value="GERENCIA RECEPTORA")
    ws.cell(row=r2, column=4, value=acta.gerencia_receptora or "")
    current_row += 1

    # Fila Cliente Funcional / Unidad Funcional
    r3 = current_row
    ws.cell(row=r3, column=1, value="NOMBRE CLIENTE FUNCIONAL")
    ws.cell(row=r3, column=2, value=acta.nombre_cliente_funcional or "")
    ws.cell(row=r3, column=3, value="UNIDAD FUNCIONAL")
    ws.cell(row=r3, column=4, value=acta.unidad_funcional or "")
    current_row += 1

    for row_idx in (r1, r2, r3):
        for col_idx in range(1, 5):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            if col_idx in (1, 3):
                c.font = font_label
                c.fill = label_fill
                c.alignment = Alignment(vertical='center', horizontal='left' if col_idx == 1 else 'center')
            else:
                c.font = font_value
                c.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)

    # Helper para bloques de texto (Descripción, Entregables, Observaciones, Detalles)
    def agregar_bloque_texto(titulo, contenido, min_height=45):
        nonlocal current_row
        current_row += 1
        ws.cell(row=current_row, column=1, value=titulo)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        apply_style_range(1, current_row, 4, current_row, font=font_section_header, fill=header_fill, alignment=Alignment(horizontal='center', vertical='center'), border=thin_border)
        current_row += 1

        content_row = current_row
        ws.cell(row=content_row, column=1, value=contenido or "")
        ws.merge_cells(start_row=content_row, start_column=1, end_row=content_row, end_column=4)
        apply_style_range(1, content_row, 4, content_row, font=font_value, alignment=Alignment(vertical='top', horizontal='left', wrap_text=True), border=thin_border)
        ws.row_dimensions[content_row].height = min_height
        current_row += 1

    agregar_bloque_texto("DESCRIPCIÓN", acta.descripcion, min_height=55)
    agregar_bloque_texto("ENTREGABLES", acta.entregables, min_height=50)
    agregar_bloque_texto("OBSERVACIONES", acta.observaciones, min_height=45)
    agregar_bloque_texto("DETALLES", acta.detalles, min_height=45)

    # SECCIÓN: FIRMAS AUTORIZADAS
    current_row += 1
    ws.cell(row=current_row, column=1, value="FIRMAS AUTORIZADAS\nGGDSTI / GGPPTI / GGLITI")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    apply_style_range(1, current_row, 4, current_row, font=font_section_header, fill=header_fill, alignment=Alignment(horizontal='center', vertical='center', wrap_text=True), border=thin_border)
    ws.row_dimensions[current_row].height = 28
    current_row += 1

    # 8 bloques de firmas en pares (izquierda: col 1-2, derecha: col 3-4)
    firmantes = [
        (acta.firmante_1_nombre, acta.firmante_1_unidad),
        (acta.firmante_2_nombre, acta.firmante_2_unidad),
        (acta.firmante_3_nombre, acta.firmante_3_unidad),
        (acta.firmante_4_nombre, acta.firmante_4_unidad),
        (acta.firmante_5_nombre, acta.firmante_5_unidad),
        (acta.firmante_6_nombre, acta.firmante_6_unidad),
        (acta.firmante_7_nombre, acta.firmante_7_unidad),
        (acta.firmante_8_nombre, acta.firmante_8_unidad),
    ]

    for idx in range(0, 8, 2):
        f_izq_nombre, f_izq_unidad = firmantes[idx]
        f_der_nombre, f_der_unidad = firmantes[idx + 1]

        start_r = current_row
        # Fila Aprobado por
        ws.cell(row=start_r, column=1, value="Aprobado por: _____/_____/_____")
        ws.merge_cells(start_row=start_r, start_column=1, end_row=start_r, end_column=2)

        ws.cell(row=start_r, column=3, value="Aprobado por: _____/_____/_____")
        ws.merge_cells(start_row=start_r, start_column=3, end_row=start_r, end_column=4)

        # Fila Nombre y Apellido
        ws.cell(row=start_r + 1, column=1, value=f"Nombre y Apellido: {f_izq_nombre or ''}")
        ws.merge_cells(start_row=start_r + 1, start_column=1, end_row=start_r + 1, end_column=2)

        ws.cell(row=start_r + 1, column=3, value=f"Nombre y Apellido: {f_der_nombre or ''}")
        ws.merge_cells(start_row=start_r + 1, start_column=3, end_row=start_r + 1, end_column=4)

        # Fila Firma
        ws.cell(row=start_r + 2, column=1, value="Firma: ____________________")
        ws.merge_cells(start_row=start_r + 2, start_column=1, end_row=start_r + 2, end_column=2)

        ws.cell(row=start_r + 2, column=3, value="Firma: ____________________")
        ws.merge_cells(start_row=start_r + 2, start_column=3, end_row=start_r + 2, end_column=4)

        # Fila Unidad
        ws.cell(row=start_r + 3, column=1, value=f"Unidad: {f_izq_unidad or ''}")
        ws.merge_cells(start_row=start_r + 3, start_column=1, end_row=start_r + 3, end_column=2)

        ws.cell(row=start_r + 3, column=3, value=f"Unidad: {f_der_unidad or ''}")
        ws.merge_cells(start_row=start_r + 3, start_column=3, end_row=start_r + 3, end_column=4)

        # Estilo para cada bloque (izq y der)
        apply_style_range(1, start_r, 2, start_r + 3, font=font_value, alignment=Alignment(vertical='center', horizontal='left'), border=thin_border)
        apply_style_range(3, start_r, 4, start_r + 3, font=font_value, alignment=Alignment(vertical='center', horizontal='left'), border=thin_border)

        current_row += 5  # 4 filas + 1 fila espaciadora

    # ESTADO OFICIAL DE APROBACIÓN
    if acta.estado == 'APROBADO':
        aprob_nom = (acta.aprobado_por.get_full_name() or acta.aprobado_por.username) if acta.aprobado_por else ""
        fecha_aprob = acta.fecha_aprobacion.strftime("%d/%m/%Y %I:%M %p") if acta.fecha_aprobacion else ""
        msg = f"✅ PROYECTO OFICIALMENTE APROBADO POR EL JEFE DE PROYECTOS ({aprob_nom}) EL {fecha_aprob}"

        ws.cell(row=current_row, column=1, value=msg)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        apply_style_range(1, current_row, 4, current_row, font=Font(name='Arial', size=9, bold=True, color='166534'), fill=approved_fill, alignment=Alignment(horizontal='center', vertical='center', wrap_text=True), border=thin_border)
        ws.row_dimensions[current_row].height = 25
    else:
        msg = "ESTADO DEL ACTA: PENDIENTE DE APROBACIÓN POR EL JEFE DE PROYECTOS"
        ws.cell(row=current_row, column=1, value=msg)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        apply_style_range(1, current_row, 4, current_row, font=Font(name='Arial', size=9, bold=True, color='854D0E'), fill=PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid'), alignment=Alignment(horizontal='center', vertical='center'), border=thin_border)
        ws.row_dimensions[current_row].height = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
